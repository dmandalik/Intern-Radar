"""Excel export support."""

from __future__ import annotations

from collections import Counter
from pathlib import Path
import json
import shutil
import subprocess
from typing import Any

from internradar.export import ExportJobView

SHEET_ORDER = [
    "Summary",
    "Open Jobs",
    "Hidden Gems",
    "Coming Soon",
    "Saved",
    "Applied",
    "Closed",
    "All Jobs",
]


def export_xlsx(
    records: list[ExportJobView],
    output_path: Path,
    *,
    metadata: dict[str, Any],
) -> Path:
    """Write practical multi-sheet XLSX export."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError:
        return _export_xlsx_with_venv(records, output_path, metadata=metadata)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_sheet = workbook.create_sheet("Summary")
    _write_summary_sheet(summary_sheet, records, metadata)

    sheet_rows = {
        "Open Jobs": [record for record in records if record.status in {"open", "likely_open"}],
        "Hidden Gems": [record for record in records if record.job.scores.hidden_gem_score >= 70.0],
        "Coming Soon": [record for record in records if record.status == "coming_soon"],
        "Saved": [record for record in records if record.saved],
        "Applied": [record for record in records if record.applied],
        "Closed": [record for record in records if record.status == "closed"],
        "All Jobs": records,
    }

    header_fill = PatternFill(fill_type="solid", fgColor="1E293B")
    header_font = Font(color="F8FAFC", bold=True)
    for sheet_name in SHEET_ORDER[1:]:
        sheet = workbook.create_sheet(sheet_name)
        rows = sheet_rows[sheet_name]
        _write_job_sheet(sheet, rows)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        _autosize_columns(sheet, get_column_letter)

    workbook.save(output_path)
    return output_path


def _write_summary_sheet(sheet: Any, records: list[ExportJobView], metadata: dict[str, Any]) -> None:
    status_counts = Counter(record.status for record in records)
    role_counts = Counter(record.role_family for record in records)
    company_counts = Counter(record.company for record in records)

    rows = [
        ("Generated at", metadata.get("generated_at", "")),
        ("Pack", metadata.get("pack", "")),
        ("Total jobs", len(records)),
        ("Open jobs", status_counts.get("open", 0)),
        ("Likely open jobs", status_counts.get("likely_open", 0)),
        ("Coming soon jobs", status_counts.get("coming_soon", 0)),
        ("Closed jobs", status_counts.get("closed", 0)),
        ("Unknown jobs", status_counts.get("unknown", 0)),
        ("Hidden gems count", sum(1 for record in records if record.job.scores.hidden_gem_score >= 70.0)),
        ("Saved count", sum(1 for record in records if record.saved)),
        ("Applied count", sum(1 for record in records if record.applied)),
        ("Top role families", ", ".join(f"{name} ({count})" for name, count in role_counts.most_common(5))),
        ("Top companies by count", ", ".join(f"{name} ({count})" for name, count in company_counts.most_common(5))),
    ]
    for index, (label, value) in enumerate(rows, start=1):
        sheet.cell(row=index, column=1, value=label)
        sheet.cell(row=index, column=2, value=value)


def _write_job_sheet(sheet: Any, records: list[ExportJobView]) -> None:
    headers = records[0].as_flat_row().keys() if records else _default_headers()
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column_index, value=header)

    if not records:
        sheet.cell(row=2, column=1, value="No jobs matched this section.")
        return

    for row_index, record in enumerate(records, start=2):
        row = record.as_flat_row()
        for column_index, header in enumerate(headers, start=1):
            value = row[header]
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if header in {"Apply URL", "Source URL"} and value:
                cell.hyperlink = value
                cell.style = "Hyperlink"


def _autosize_columns(sheet: Any, get_column_letter: Any) -> None:
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 42)


def _default_headers() -> list[str]:
    return list(ExportJobView(rank=1, job=_empty_job()).as_flat_row().keys())


def _empty_job() -> Any:
    from datetime import UTC, datetime

    from internradar.core.models import ClassifiedRole, EligibilityInfo, Job, JobScores, JobStatusInfo

    now = datetime.now(UTC)
    return Job(
        id="empty",
        company_id="empty",
        company_name="",
        title="",
        description=None,
        apply_url="",
        source_url="",
        source_type="",
        role=ClassifiedRole(role_family="unknown", confidence=0.0),
        season=None,
        year=None,
        locations=[],
        remote_type=None,
        status=JobStatusInfo(status="unknown", confidence=0.0, evidence=[], checked_at=now),
        eligibility=EligibilityInfo(),
        scores=JobScores(),
        prestige_tier=None,
        tags=[],
        first_seen=now,
        last_seen=now,
        last_verified=now,
        content_hash=None,
    )


def _export_xlsx_with_venv(
    records: list[ExportJobView],
    output_path: Path,
    *,
    metadata: dict[str, Any],
) -> Path:
    python_command = shutil.which("python3")
    if python_command is None:
        raise ModuleNotFoundError("openpyxl is not available and python3 could not be located for fallback export.")

    payload = {
        "output_path": str(output_path),
        "metadata": metadata,
        "records": [_record_payload(record) for record in records],
    }
    script = """
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

payload = json.loads(sys.argv[1])
output_path = Path(payload["output_path"])
output_path.parent.mkdir(parents=True, exist_ok=True)
records = payload["records"]
metadata = payload["metadata"]

def headers():
    return [
        "Rank","Company","Title","Role Family","Role Subtype","Status","Season","Year","Location",
        "Remote Type","Apply URL","Source URL","Prestige Tier","Opportunity Score","Role Fit Score",
        "Technical Depth Score","Eligibility Score","Hidden Gem Score","Freshness Score","First Seen",
        "Last Seen","Last Verified","Application Status","Eligibility Summary","Why Ranked","Notes"
    ]

def summary_rows():
    summary = metadata["summary"]
    return [
        ("Generated at", metadata.get("generated_at", "")),
        ("Pack", metadata.get("pack", "")),
        ("Total jobs", len(records)),
        ("Open jobs", summary.get("open", 0)),
        ("Likely open jobs", summary.get("likely_open", 0)),
        ("Coming soon jobs", summary.get("coming_soon", 0)),
        ("Closed jobs", summary.get("closed", 0)),
        ("Unknown jobs", summary.get("unknown", 0)),
        ("Hidden gems count", summary.get("hidden_gems", 0)),
        ("Saved count", summary.get("saved", 0)),
        ("Applied count", summary.get("applied", 0)),
    ]

def autosize(sheet):
    for column_cells in sheet.columns:
        max_len = 0
        col = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        sheet.column_dimensions[col].width = min(max(max_len + 2, 12), 42)

wb = Workbook()
wb.remove(wb.active)
summary_sheet = wb.create_sheet("Summary")
for row_index, (label, value) in enumerate(summary_rows(), start=1):
    summary_sheet.cell(row=row_index, column=1, value=label)
    summary_sheet.cell(row=row_index, column=2, value=value)

sheet_groups = {
    "Open Jobs": [record for record in records if record["Status"] in {"open", "likely_open"}],
    "Hidden Gems": [record for record in records if float(record["Hidden Gem Score"]) >= 70.0],
    "Coming Soon": [record for record in records if record["Status"] == "coming_soon"],
    "Saved": [record for record in records if record["saved"]],
    "Applied": [record for record in records if record["applied"]],
    "Closed": [record for record in records if record["Status"] == "closed"],
    "All Jobs": records,
}
header_fill = PatternFill(fill_type="solid", fgColor="1E293B")
header_font = Font(color="F8FAFC", bold=True)
for name, rows in sheet_groups.items():
    sheet = wb.create_sheet(name)
    for col_index, header in enumerate(headers(), start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
    if not rows:
        sheet.cell(row=2, column=1, value="No jobs matched this section.")
    else:
        for row_index, row in enumerate(rows, start=2):
            for col_index, header in enumerate(headers(), start=1):
                value = row.get(header, "")
                cell = sheet.cell(row=row_index, column=col_index, value=value)
                if header in {"Apply URL", "Source URL"} and value:
                    cell.hyperlink = value
                    cell.style = "Hyperlink"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    autosize(sheet)
wb.save(output_path)
"""
    subprocess.run(
        [python_command, "-c", script, json.dumps(payload)],
        check=True,
        capture_output=True,
        text=True,
    )
    return output_path


def _record_payload(record: ExportJobView) -> dict[str, Any]:
    row = record.as_flat_row()
    row["saved"] = record.saved
    row["applied"] = record.applied
    return row


__all__ = ["SHEET_ORDER", "export_xlsx"]
